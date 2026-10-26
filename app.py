"""stuffid - stage 1: inspect the product_good zip files from Google Drive.

Endpoints (all except / require ?token=ADMIN_TOKEN):
  /                      health
  /files                 list configured Drive files
  /inspect?part=7        download one part, list zip contents, show sheets, header and sample rows
"""
import json
import os
import threading
import time
import zipfile

import requests
from flask import Flask, Response, request

app = Flask(__name__)

TOKEN = os.environ.get("ADMIN_TOKEN", "")
WORK = "/tmp/stuffid"
os.makedirs(WORK, exist_ok=True)

# Google Drive file ids, folder "tax"
PARTS = {
    1: "1ELetseHnuh5g_sMY_sJkO0vDlSgVTmjZ",
    2: "1n8dPhscbVBIWgHsri9wQhnqYrIebyEKe",
    3: "1OhGIS6IVVJzDD8cct0PXBwc2TtZmmH_7",
    4: "1eg1WMh-e4pAgsEmn2Sqjo8IU0hQJiJzL",
    5: "1Yu2joZUKRwKB73r5rSwvHvP4H0YPQ-Wb",
    6: "1AoodP9jbSfuqwAEUXp_DHgBmqyonucXM",
    7: "1WA-WMEcXjY6_NH4uAFHYJxyS5xswiEQA",
}

_cache = {}
_lock = threading.Lock()


def as_json(obj, status=200):
    return Response(json.dumps(obj, ensure_ascii=False, indent=1, default=str),
                    status=status, mimetype="application/json; charset=utf-8")


def authorized():
    return TOKEN and request.args.get("token") == TOKEN


def download(part):
    """Download a public Drive file (link sharing on). Returns local path."""
    fid = PARTS[part]
    path = os.path.join(WORK, f"part_{part}.zip")
    if os.path.exists(path) and zipfile.is_zipfile(path):
        return path
    url = "https://drive.usercontent.google.com/download"
    with requests.get(url, params={"id": fid, "export": "download", "confirm": "t"},
                      stream=True, timeout=120) as r:
        r.raise_for_status()
        ctype = r.headers.get("content-type", "")
        if "text/html" in ctype:
            snippet = r.raw.read(600).decode("utf-8", "replace")
            raise RuntimeError(f"Drive returned HTML instead of the file (is link sharing on?): {snippet[:300]}")
        tmp = path + ".part"
        with open(tmp, "wb") as f:
            for chunk in r.iter_content(1 << 20):
                f.write(chunk)
    os.replace(tmp, path)
    return path


def inspect_xlsx(fobj, max_rows=6):
    from openpyxl import load_workbook
    wb = load_workbook(fobj, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(list(row))
            if i + 1 >= max_rows:
                break
        sheets.append({"sheet": ws.title, "max_row": ws.max_row,
                       "max_column": ws.max_column, "rows": rows})
    wb.close()
    return sheets


def inspect_csv(fobj, max_rows=6):
    import io
    raw = fobj.read(200_000)
    for enc in ("utf-8-sig", "cp1256", "utf-16"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = text.splitlines()[:max_rows]
    return [{"encoding": enc, "lines": lines}]


def inspect_part(part):
    t0 = time.time()
    path = download(part)
    out = {"part": part, "zip_bytes": os.path.getsize(path), "download_s": round(time.time() - t0, 1),
           "entries": []}
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        out["entry_count"] = len(infos)
        for zi in infos[:50]:
            e = {"name": zi.filename, "size": zi.file_size, "compressed": zi.compress_size}
            out["entries"].append(e)
        # inspect the first two data files in detail
        detailed = 0
        for zi in infos:
            name = zi.filename.lower()
            if detailed >= 2:
                break
            try:
                if name.endswith((".xlsx", ".xlsm")):
                    with z.open(zi) as f:
                        import io
                        data = io.BytesIO(f.read())
                    out.setdefault("details", []).append({"file": zi.filename, "sheets": inspect_xlsx(data)})
                    detailed += 1
                elif name.endswith((".csv", ".txt")):
                    with z.open(zi) as f:
                        out.setdefault("details", []).append({"file": zi.filename, "csv": inspect_csv(f)})
                    detailed += 1
            except Exception as ex:  # keep going, report error
                out.setdefault("details", []).append({"file": zi.filename, "error": repr(ex)})
                detailed += 1
    out["total_s"] = round(time.time() - t0, 1)
    return out


@app.get("/")
def health():
    return as_json({"ok": True, "service": "stuffid", "stage": "inspect"})


@app.get("/files")
def files():
    if not authorized():
        return as_json({"error": "unauthorized"}, 401)
    return as_json(PARTS)


@app.get("/inspect")
def inspect():
    if not authorized():
        return as_json({"error": "unauthorized"}, 401)
    part = int(request.args.get("part", "7"))
    if part not in PARTS:
        return as_json({"error": "unknown part"}, 400)
    with _lock:
        if part in _cache and not request.args.get("refresh"):
            return as_json(_cache[part])
    try:
        res = inspect_part(part)
    except Exception as ex:
        return as_json({"part": part, "error": repr(ex)}, 500)
    with _lock:
        _cache[part] = res
    return as_json(res)


def log_result(tag, obj):
    """Write a result to stdout in chunks so it can be read back from Render logs."""
    text = json.dumps(obj, ensure_ascii=False, default=str)
    size = 2500
    chunks = [text[i:i + size] for i in range(0, len(text), size)] or [""]
    for n, c in enumerate(chunks, 1):
        print(f"RESULT {tag} {n}/{len(chunks)} {c}", flush=True)


def keepalive():
    """Free Render services sleep after ~15 min without traffic; ping ourselves while a job runs."""
    url = os.environ.get("RENDER_EXTERNAL_URL")
    if url:
        try:
            requests.get(url, timeout=20)
        except Exception:
            pass


def stats_part(part, ids_out):
    """Stream the CSV inside one zip and collect summary statistics."""
    import csv
    import io
    from array import array
    from collections import Counter
    t0 = time.time()
    path = download(part)
    st = {"part": part, "rows": 0, "bad_rows": 0, "bad_id": 0, "desc_len_sum": 0, "desc_len_max": 0,
          "csv_bytes": 0, "Type": Counter(), "Vat": Counter(), "Taxable": Counter(),
          "has_expiration": 0, "has_pricing": 0, "id_prefix": Counter(),
          "run_year": Counter(), "examples_bad": []}
    ids = array("q")
    last_ping = time.time()
    with zipfile.ZipFile(path) as z:
        zi = [i for i in z.infolist() if i.filename.lower().endswith(".csv")][0]
        st["csv_bytes"] = zi.file_size
        with z.open(zi) as raw:
            reader = csv.reader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline=""))
            header = next(reader)
            st["header"] = header
            for row in reader:
                if len(row) != 10:
                    st["bad_rows"] += 1
                    if len(st["examples_bad"]) < 3:
                        st["examples_bad"].append(row[:12])
                    continue
                st["rows"] += 1
                rid = row[0].strip()
                if rid.isdigit() and len(rid) == 13:
                    ids.append(int(rid))
                    st["id_prefix"][rid[:2]] += 1
                else:
                    st["bad_id"] += 1
                d = len(row[1])
                st["desc_len_sum"] += d
                if d > st["desc_len_max"]:
                    st["desc_len_max"] = d
                st["Vat"][row[2]] += 1
                st["Taxable"][row[3]] += 1
                st["run_year"][row[4][:4]] += 1
                if row[5]:
                    st["has_expiration"] += 1
                st["Type"][row[8]] += 1
                if row[9]:
                    st["has_pricing"] += 1
                if st["rows"] % 500000 == 0:
                    print(f"PROGRESS part {part} rows {st['rows']} t={round(time.time() - t0)}s", flush=True)
                if time.time() - last_ping > 240:
                    keepalive()
                    last_ping = time.time()
    import numpy as np
    st["unique_ids_in_part"] = int(np.unique(np.frombuffer(ids, dtype=np.int64)).size)
    ids_out.extend(ids)
    del ids
    for k in ("Type", "Vat", "Taxable", "id_prefix", "run_year"):
        st[k] = dict(st[k].most_common(25))
    st["seconds"] = round(time.time() - t0)
    try:
        os.remove(path)  # free the small /tmp disk
    except OSError:
        pass
    return st


def stats_all(parts):
    from array import array
    all_ids = array("q")
    summary = {"parts": {}}
    for p in parts:
        try:
            st = stats_part(p, all_ids)
            log_result(f"stats:{p}", st)
            summary["parts"][p] = st["rows"]
        except Exception as ex:
            log_result(f"stats:{p}", {"error": repr(ex)})
        keepalive()
    import numpy as np
    a = np.frombuffer(all_ids, dtype=np.int64).copy()
    a.sort()
    dup_mask = a[1:] == a[:-1]
    summary["total_rows_with_valid_id"] = int(a.size)
    summary["duplicate_id_rows"] = int(dup_mask.sum())
    summary["unique_ids"] = int(a.size - dup_mask.sum())
    summary["sample_duplicate_ids"] = [int(x) for x in np.unique(a[1:][dup_mask])[:10]]
    log_result("stats:total", summary)


def iter_rows(part):
    """Yield CSV rows (lists of 10 fields) of one part."""
    import csv
    import io
    path = download(part)
    with zipfile.ZipFile(path) as z:
        zi = [i for i in z.infolist() if i.filename.lower().endswith(".csv")][0]
        with z.open(zi) as raw:
            reader = csv.reader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline=""))
            next(reader)
            for row in reader:
                if len(row) == 10:
                    yield row


def match_products(names, top=6, parts=None):
    """Score every stuff-ID row against each product name (token overlap weighted by IDF)."""
    import heapq
    import math
    from collections import defaultdict
    from matcher import tokens

    prods = [list(dict.fromkeys(tokens(n))) for n in names]
    vocab = defaultdict(list)
    for i, toks in enumerate(prods):
        for t in toks:
            vocab[t].append(i)

    t0 = time.time()
    df = defaultdict(int)
    n_rows = 0
    for row in iter_rows(7):  # ~82k rows: a large enough sample to estimate token rarity
        n_rows += 1
        for t in set(tokens(row[1])):
            if t in vocab:
                df[t] += 1
    print(f"PROGRESS df sample rows={n_rows} t={round(time.time() - t0)}s", flush=True)
    idf = {t: max(0.05, math.log((n_rows + 1) / (df.get(t, 0) + 1))) for t in vocab}
    total = [sum(idf[t] for t in toks) or 1.0 for toks in prods]
    pset = [set(toks) for toks in prods]

    spec = [({}, []) for _ in names]   # specific IDs: score = share of product name covered
    gen = [({}, []) for _ in names]    # general IDs: score = share of ID description covered by product

    def offer(slot, score, row, cap):
        d, h = slot
        rid = row[0]
        prev = d.get(rid)
        if prev is not None:
            if score > prev[0] or (score == prev[0] and row[7] > prev[1][7]):
                d[rid] = (max(score, prev[0]), row)
            return
        if len(h) < cap:
            heapq.heappush(h, (score, rid))
            d[rid] = (score, row)
        elif score > h[0][0]:
            _, old = heapq.heapreplace(h, (score, rid))
            d.pop(old, None)
            d[rid] = (score, row)

    for p in (parts or sorted(PARTS)):
        for row in iter_rows(p):
            toks = set(tokens(row[1]))
            is_general = "عمومی" in row[8]
            acc = defaultdict(float)
            for t in toks:
                for i in vocab.get(t, ()):
                    acc[i] += idf[t]
            for i, s in acc.items():
                if is_general:
                    shared = sum(1 for t in toks if t in pset[i])
                    score = s / (s + 2.5 * (len(toks) - shared)) if s > 0 else 0.0
                    if score >= 0.3:
                        offer(gen[i], score, row, 4)
                else:
                    score = s / total[i]
                    if score >= 0.25:
                        offer(spec[i], score, row, top)
        print(f"PROGRESS score part {p} t={round(time.time() - t0)}s", flush=True)
        keepalive()

    def fmt(slot):
        return [{"score": round(sc, 3), "id": r[0], "desc": r[1], "vat": r[2], "taxable": r[3],
                 "exp": r[5], "edited": r[7], "type": r[8]}
                for sc, r in sorted(slot[0].values(), key=lambda x: -x[0])]

    return [{"product": name, "tokens": prods[i], "specific": fmt(spec[i]), "general": fmt(gen[i])}
            for i, name in enumerate(names)]


def export_general():
    """Log every general (عمومی) stuff ID once, latest edit wins, as compact pipe-separated lines."""
    best = {}
    for p in sorted(PARTS):
        for row in iter_rows(p):
            if "عمومی" not in row[8]:
                continue
            prev = best.get(row[0])
            if prev is None or row[7] > prev[7]:
                best[row[0]] = row
        print(f"PROGRESS general part {p} unique={len(best)}", flush=True)
    origin = {"شناسه عمومی وارداتی": "I", "شناسه عمومی تولید داخل": "D"}
    lines = []
    for rid in sorted(best):
        r = best[rid]
        desc = r[1].replace("|", "/").replace("\n", " ")
        lines.append("|".join([rid, origin.get(r[8], r[8]), r[2], r[3], r[5], r[7], desc]))
    # pack lines into ~3000-char log messages (lines never split)
    buf, n = [], 0
    for ln in lines:
        if buf and sum(len(x) + 1 for x in buf) + len(ln) > 3000:
            n += 1
            print("GEN " + str(n) + " " + "\n".join(buf).replace("\n", " ¶ "), flush=True)
            buf = []
        buf.append(ln)
    if buf:
        n += 1
        print("GEN " + str(n) + " " + "\n".join(buf).replace("\n", " ¶ "), flush=True)
    print(f"GENDONE chunks={n} rows={len(lines)}", flush=True)


def brand_origin(parts):
    """Count, for each brand alias, how many specific IDs mention it, split by imported/domestic."""
    from collections import defaultdict
    from matcher import normalize
    here = os.path.dirname(os.path.abspath(__file__))
    with open(os.path.join(here, "brand_aliases.json"), encoding="utf-8") as f:
        aliases = set(json.load(f))
    maxn = max(len(a.split()) for a in aliases)
    counts = defaultdict(lambda: [0, 0])  # [imported, domestic]
    t0 = time.time()
    for p in parts:
        for row in iter_rows(p):
            t = row[8]
            if "اختصاصی" not in t:
                continue
            k = 0 if "وارداتی" in t else 1
            w = normalize(row[1]).split()
            seen = set()
            for i in range(len(w)):
                for j in range(1, maxn + 1):
                    if i + j > len(w):
                        break
                    g = " ".join(w[i:i + j])
                    if g in aliases:
                        seen.add(g)
            for g in seen:
                counts[g][k] += 1
        print(f"PROGRESS brand part {p} t={round(time.time() - t0)}s", flush=True)
        keepalive()
    log_result("brand_origin:" + ",".join(map(str, parts)), dict(counts))


def run_job():
    job = os.environ.get("JOB", "").strip()
    if not job:
        return
    print(f"JOB start {job}", flush=True)
    for step in job.split(";"):
        kind, _, arg = step.partition(":")
        try:
            if kind == "inspect":
                log_result(step, inspect_part(int(arg)))
            elif kind == "match_sample":
                here = os.path.dirname(os.path.abspath(__file__))
                with open(os.path.join(here, "sample_products.txt"), encoding="utf-8") as f:
                    names = [x.strip() for x in f if x.strip()]
                parts = [int(x) for x in arg.split(",")] if arg else None
                for n, res in enumerate(match_products(names, parts=parts), 1):
                    log_result(f"match:{n}", res)
            elif kind == "brand_origin":
                brand_origin([int(x) for x in arg.split(",")] if arg else sorted(PARTS))
            elif kind == "export_general":
                export_general()
            elif kind == "stats":
                parts = [int(x) for x in arg.split(",")] if arg else sorted(PARTS)
                stats_all(parts)
        except Exception as ex:
            log_result(step, {"error": repr(ex)})
        keepalive()
    print(f"JOB done {job}", flush=True)


threading.Thread(target=run_job, daemon=True).start()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 10000)))
