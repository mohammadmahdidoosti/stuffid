"""Review workbook: every product with its suggested general stuff ID.

usage: build_review_xlsx.py products.json general_ids.tsv brands.json out.xlsx
"""
import csv
import json
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import rules

products_f, general_f, brands_f, out_f = sys.argv[1:5]
names = json.load(open(products_f, encoding="utf-8"))
gen = defaultdict(dict)
for r in list(csv.reader(open(general_f, encoding="utf-8"), delimiter="\t"))[1:]:
    gen[r[6]][r[1]] = r
brands = json.load(open(brands_f, encoding="utf-8"))
DEFAULT_ORIGIN = "D"


def origin(n):
    s = " " + rules.norm(n) + " "
    for a, o in brands:
        if " " + a + " " in s:
            return o, "برند"
    return DEFAULT_ORIGIN, "پیش‌فرض (برند ناشناخته)"


rows = []
for n in names:
    desc, _ = rules.classify(n)
    if not desc:
        rows.append([n, "", "", "", "", None, "بدون قاعده — دستی انتخاب شود"])
        continue
    o, src = origin(n)
    variants = gen[desc]
    r = variants.get(o) or next(iter(variants.values()))
    note = "پیشنهاد" if r[1] == o else "پیشنهاد — فقط نسخه " + ("وارداتی" if r[1] == "I" else "داخلی") + " وجود دارد"
    rows.append([n, r[0], desc, "وارداتی" if o == "I" else "تولید داخل", src, float(r[2]), note])

F = "Arial"
wb = Workbook()
ws = wb.active
ws.title = "نتیجه"
ws.sheet_view.rightToLeft = True
head = ["نام محصول", "شناسه پیشنهادی", "شرح شناسه عمومی", "مبدا", "تشخیص مبدا", "نرخ مالیات (%)", "وضعیت"]
ws.append(head)
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(name=F, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F)
    row[1].number_format = "@"
    if row[6].value and row[6].value.startswith("بدون"):
        row[6].font = Font(name=F, color="C00000")
for i, w in enumerate([70, 17, 60, 12, 22, 14, 34], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
last = len(rows) + 1
t = Table(displayName="Results", ref=f"A1:G{last}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
ws.add_table(t)

s = wb.create_sheet("خلاصه", 0)
s.sheet_view.rightToLeft = True
R = f"'نتیجه'!$B$2:$B${last}"
lines = [
    ("کل محصولات", f"=COUNTA('نتیجه'!$A$2:$A${last})"),
    ("دارای شناسه پیشنهادی", f'=COUNTIF({R},"?*")'),
    ("بدون قاعده (انتخاب دستی)", "=B2-B3"),
    ("درصد پوشش", "=IF(B2=0,0,B3/B2)"),
    ("وارداتی", f"=COUNTIFS('نتیجه'!$D$2:$D${last},\"وارداتی\")"),
    ("تولید داخل", f"=COUNTIFS('نتیجه'!$D$2:$D${last},\"تولید داخل\")"),
    ("مبدا از روی برند", f"=COUNTIFS('نتیجه'!$E$2:$E${last},\"برند\")"),
    ("مبدا پیش‌فرض (برند ناشناخته)", f"=COUNTIFS('نتیجه'!$E$2:$E${last},\"پیش‌فرض*\")"),
]
s.append(["شاخص", "مقدار"])
for a, b in lines:
    s.append([a, b])
s["B5"].number_format = "0.0%"
for row in s.iter_rows(min_row=1, max_row=len(lines) + 1):
    for c in row:
        c.font = Font(name=F, bold=(c.row == 1))
s.column_dimensions["A"].width = 34
s.column_dimensions["B"].width = 14
notes = [
    "راهنما:",
    "• شناسه‌ها «شناسه عمومی» هستند، بر اساس نوع کالا (از نام محصول) و مبدا (از برند).",
    "• اگر برند ناشناخته بود، «تولید داخل» فرض شده است (قابل تغییر در تنظیمات افزونه).",
    "• طبق ضوابط، اگر شناسه اختصاصی کالا را دارید (از صورتحساب خرید)، آن اولویت دارد.",
    "• ردیف‌های «بدون قاعده» را در افزونه با جستجو انتخاب کنید.",
    "منبع: فهرست شناسه‌های عمومی از فایل‌های product_good (۲۵ مهر ۱۴۰۵).",
]
for i, line in enumerate(notes, len(lines) + 4):
    s.cell(row=i, column=1, value=line).font = Font(name=F, bold=(line == "راهنما:"))
wb.save(out_f)
print("rows", len(rows))
